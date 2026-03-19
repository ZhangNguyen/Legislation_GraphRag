from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import List

from pypdf import PdfReader

try:
    from docx import Document  # python-docx
except Exception:
    Document = None


def read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def read_pdf_pypdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: List[str] = []
    for page in reader.pages:
        t = page.extract_text() or ""
        parts.append(t)
    return "\n".join(parts)


def read_docx(path: Path) -> str:
    if Document is None:
        raise RuntimeError(
            "Missing dependency 'python-docx'. Install: pip install python-docx"
        )
    doc = Document(str(path))
    parts: List[str] = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            parts.append(txt)

    # đọc thêm text trong bảng nếu có
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_txt = cell.text.strip()
                if cell_txt:
                    row_text.append(cell_txt)
            if row_text:
                parts.append(" | ".join(row_text))

    return "\n".join(parts)


def _find_soffice() -> str | None:
    candidates = [
        shutil.which("soffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for c in candidates:
        if c and Path(c).exists():
            return c
    return None


def convert_doc_to_docx_with_soffice(path: Path) -> Path:
    soffice = _find_soffice()
    if not soffice:
        raise RuntimeError(
            "Không tìm thấy LibreOffice (soffice). "
            "Hãy cài LibreOffice hoặc chuyển .doc sang .docx trước."
        )

    tmpdir = Path(tempfile.mkdtemp(prefix="doc_convert_"))
    cmd = [
        soffice,
        "--headless",
        "--convert-to",
        "docx",
        "--outdir",
        str(tmpdir),
        str(path),
    ]
    result = subprocess.run(cmd, capture_output=True, text=True)

    if result.returncode != 0:
        raise RuntimeError(
            f"LibreOffice convert failed for {path.name}\n"
            f"STDOUT: {result.stdout}\nSTDERR: {result.stderr}"
        )

    out_docx = tmpdir / f"{path.stem}.docx"
    if not out_docx.exists():
        raise RuntimeError(f"Converted file not found: {out_docx}")
    return out_docx


def read_doc(path: Path) -> str:
    # Cách ổn định nhất đa nền tảng: convert .doc -> .docx bằng LibreOffice rồi đọc bằng python-docx
    converted = convert_doc_to_docx_with_soffice(path)
    return read_docx(converted)


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return text.strip()


def load_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return normalize_text(read_txt(path))
    if suffix == ".pdf":
        return normalize_text(read_pdf_pypdf(path))
    if suffix == ".docx":
        return normalize_text(read_docx(path))
    if suffix == ".doc":
        return normalize_text(read_doc(path))
    raise ValueError(f"Unsupported file type: {suffix}")